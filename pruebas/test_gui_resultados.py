"""Creado con Claude. Ciclo completo en la interfaz: generar → fijar → regenerar → exportar."""

from pathlib import Path

import openpyxl
import pytest
from PySide6.QtCore import QSettings, QTimer
from pytestqt.qtbot import QtBot

from timetabling.dominio.modelos import EstadoSolucion
from timetabling.ejemplo import crear_proyecto_ejemplo
from timetabling.gui.app import VentanaPrincipal
from timetabling.gui.pantallas.generar import PantallaGenerar
from timetabling.gui.pantallas.resultados import PantallaResultados
from timetabling.gui.widgets.vista_horario import ICONO_FIJA
from timetabling.nucleo.verificador import verificar
from timetabling.persistencia.repositorio import Repositorio


@pytest.fixture
def ventana(qtbot: QtBot, tmp_path: Path) -> VentanaPrincipal:
    ajustes = QSettings(str(tmp_path / "ajustes.ini"), QSettings.Format.IniFormat)
    v = VentanaPrincipal(ajustes)
    qtbot.addWidget(v)
    v.show()
    ruta = tmp_path / "escuela.horario"
    crear_proyecto_ejemplo(ruta)
    v.abrir_proyecto(ruta)
    return v


def _generar(ventana: VentanaPrincipal, qtbot: QtBot, limite: float) -> PantallaGenerar:
    generar = ventana.mostrar("generar")
    assert isinstance(generar, PantallaGenerar)
    assert generar.generar(tiempo_limite_s=limite)
    return generar


def _esperar(qtbot: QtBot, generar: PantallaGenerar, timeout_s: float = 90) -> None:
    qtbot.waitUntil(lambda: not generar.en_curso, timeout=int(timeout_s * 1000))


def test_ciclo_generar_fijar_regenerar_exportar(
    ventana: VentanaPrincipal, qtbot: QtBot, tmp_path: Path
) -> None:
    # La interfaz sigue procesando eventos mientras el solver trabaja.
    latidos: list[int] = []
    reloj = QTimer()
    reloj.timeout.connect(lambda: latidos.append(1))
    reloj.start(50)

    generar = _generar(ventana, qtbot, limite=4)
    assert not ventana.pantallas["profesores"].isEnabled()
    assert not ventana.pantallas["resultados"].isEnabled()
    assert not ventana.accion_abrir.isEnabled()
    _esperar(qtbot, generar)
    reloj.stop()
    assert len(latidos) >= 20  # ~4 s a 50 ms: el bucle de eventos no se congeló
    assert ventana.pantallas["profesores"].isEnabled()
    assert ventana.accion_abrir.isEnabled()
    assert generar.ultima_solucion_id is not None
    assert "verificada y guardada" in generar.mensaje.text()

    # Resultados de la solución recién generada
    generar.boton_resultados.click()
    resultados = ventana.pantallas["resultados"]
    assert isinstance(resultados, PantallaResultados)
    assert ventana.pila.currentWidget() is resultados
    assert resultados.solucion is not None
    primera = resultados.solucion
    assert resultados.tabla_carga.rowCount() == len(ventana.repo.listar_profesores())

    # Fijar una clase desde la vista por grupo
    vista = resultados.vista_grupo
    celda = next(
        (f, c)
        for f in range(vista.rowCount())
        for c in range(vista.columnCount())
        if len(vista.asignaciones_en(f, c)) == 1
    )
    clase = vista.asignaciones_en(*celda)[0]
    resultados.fijar(clase, True)
    fijas = ventana.repo.listar_asignaciones_fijas()
    assert len(fijas) == 1 and fijas[0].profesor_id == clase.profesor_id
    assert vista.item(*celda).text().startswith(ICONO_FIJA)

    # Regenerar: la clase fija se respeta
    generar = _generar(ventana, qtbot, limite=4)
    _esperar(qtbot, generar)
    assert generar.ultima_solucion_id != primera.id
    with Repositorio.abrir(ventana.repo.ruta) as repo:
        nueva = repo.cargar_solucion(generar.ultima_solucion_id)
        assert verificar(repo.cargar_datos(), nueva.asignaciones) == []
    assert any(
        (a.grupo_id, a.asignatura_id, a.profesor_id, a.dia, a.indice)
        == (clase.grupo_id, clase.asignatura_id, clase.profesor_id, clase.dia, clase.indice)
        for a in nueva.asignaciones
    )

    # Desfijar
    ventana.ver_solucion(nueva.id)
    resultados.fijar(clase, False)
    assert ventana.repo.listar_asignaciones_fijas() == []
    assert resultados.selector.count() == 2

    # Exportar a Excel
    destino = resultados.exportar(tmp_path / "horario")
    assert destino.suffix == ".xlsx"
    libro = openpyxl.load_workbook(destino)
    assert "Resumen" in libro.sheetnames and "K3A" in libro.sheetnames


def test_cancelar_conserva_la_mejor_solucion(ventana: VentanaPrincipal, qtbot: QtBot) -> None:
    generar = _generar(ventana, qtbot, limite=120)
    qtbot.waitUntil(lambda: generar.soluciones.text() != "0", timeout=60000)
    generar.boton_cancelar.click()
    _esperar(qtbot, generar, timeout_s=30)
    assert generar.ultima_solucion_id is not None
    solucion = ventana.repo.cargar_solucion(generar.ultima_solucion_id)
    assert solucion.estado in (EstadoSolucion.CANCELADO, EstadoSolucion.OPTIMO)
    assert solucion.tiempo_s < 60


def test_infactible_muestra_diagnostico(ventana: VentanaPrincipal, qtbot: QtBot) -> None:
    repo = ventana.repo
    assert repo is not None
    titular = next(p for p in repo.listar_profesores() if p.nombre == "Ana López")
    for rango in titular.disponibilidad:
        rango.inicio = 9 * 60 + 30
    repo.guardar_profesor(titular)
    generar = _generar(ventana, qtbot, limite=10)
    _esperar(qtbot, generar)
    assert generar.caja_diagnostico.isVisible()
    assert generar.diagnostico.count() >= 1
    assert "Ana López" in "\n".join(
        generar.diagnostico.item(i).text() for i in range(generar.diagnostico.count())
    )
    assert repo.listar_soluciones() == []


def test_errores_de_validacion_impiden_generar(ventana: VentanaPrincipal) -> None:
    repo = ventana.repo
    assert repo is not None
    for profesor in repo.listar_profesores():
        repo.guardar_profesor(profesor.model_copy(update={"asignaturas": set()}))
    generar = ventana.mostrar("generar")
    assert isinstance(generar, PantallaGenerar)
    assert not generar.generar()
    assert "error(es) en los datos" in generar.mensaje.text()
    assert generar.boton_validacion.isVisibleTo(generar)
    generar.boton_validacion.click()
    assert ventana.pila.currentWidget() is ventana.pantallas["validacion"]
