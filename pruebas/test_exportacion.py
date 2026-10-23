"""Creado con Claude. Exportación a Excel: hojas, celdas combinadas e impresión (openpyxl)."""

from pathlib import Path

import openpyxl
import pytest
from openpyxl.worksheet.worksheet import Worksheet

from timetabling.dominio.cuadricula import cuadricula_grupo, cuadricula_profesor
from timetabling.dominio.modelos import DatosEscuela, Solucion
from timetabling.ejemplo import Tamano, construir_datos
from timetabling.exportacion.estilos import nombre_hoja
from timetabling.exportacion.excel import OpcionesExportacion, exportar_excel
from timetabling.nucleo.franjas import generar_franjas
from timetabling.nucleo.resolver import resolver


@pytest.fixture(scope="module")
def resuelto() -> tuple[DatosEscuela, Solucion]:
    datos = construir_datos(Tamano.PEQUENO)
    resultado = resolver(datos, tiempo_limite_s=10)
    assert resultado.asignaciones
    solucion = Solucion(
        estado=resultado.estado,
        tiempo_s=resultado.tiempo_s,
        penalizaciones=resultado.penalizaciones,
        asignaciones=resultado.asignaciones,
    )
    return datos, solucion


@pytest.fixture(scope="module")
def libro(
    resuelto: tuple[DatosEscuela, Solucion], tmp_path_factory: pytest.TempPathFactory
) -> openpyxl.Workbook:
    datos, solucion = resuelto
    ruta = tmp_path_factory.mktemp("xlsx") / "horario.xlsx"
    exportar_excel(datos, solucion, generar_franjas(datos), ruta)
    return openpyxl.load_workbook(ruta)


def _rangos(hoja: Worksheet) -> list[tuple[int, int, int, int]]:
    return [(r.min_row, r.min_col, r.max_row, r.max_col) for r in hoja.merged_cells.ranges]


def test_hojas_esperadas(libro: openpyxl.Workbook, resuelto: tuple[DatosEscuela, Solucion]) -> None:
    datos, solucion = resuelto
    nombres = libro.sheetnames
    assert nombres[0] == "Resumen"
    for grupo in datos.grupos:
        assert datos.nombre_grupo(grupo) in nombres
    con_clase = {a.profesor_id for a in solucion.asignaciones}
    for profesor in datos.profesores:
        assert (profesor.nombre in nombres) == (profesor.id in con_clase)
    assert {"patio Preescolar", "patio Primaria", "patio Secundaria", "computo Secundaria"} <= set(
        nombres
    )
    assert nombres[-2:] == ["CargaDocente", "Leyenda"]


def test_hoja_de_grupo(libro: openpyxl.Workbook, resuelto: tuple[DatosEscuela, Solucion]) -> None:
    datos, solucion = resuelto
    hoja = libro["7°A"]
    assert hoja["B2"].value == "Lunes" and hoja["F2"].value == "Viernes"
    assert hoja["A3"].value == "07:00–07:50"
    # Receso de secundaria (10:20–10:40) como fila gris combinada tras la 4.ª franja.
    fila_receso = 3 + 4
    assert (fila_receso, 1, fila_receso, 6) in _rangos(hoja)
    assert "Receso" in hoja.cell(fila_receso, 1).value
    assert hoja.cell(fila_receso, 1).fill.fgColor.rgb.endswith("D9D9D9")
    # Todas las clases del grupo aparecen con asignatura y profesor.
    grupo = next(g for g in datos.grupos if datos.nombre_grupo(g) == "7°A")
    clases = [a for a in solucion.asignaciones if a.grupo_id == grupo.id]
    textos = [c.value for fila in hoja.iter_rows(min_row=3, min_col=2) for c in fila if c.value]
    assert len([t for t in textos if "\n" in t]) == len(clases)
    profesor = datos.profesor(clases[0].profesor_id)
    assert any(profesor.nombre in t for t in textos)


def test_profesor_multinivel_con_escala_proporcional(
    libro: openpyxl.Workbook, resuelto: tuple[DatosEscuela, Solucion]
) -> None:
    datos, solucion = resuelto
    multinivel = next(
        p
        for p in datos.profesores
        if len({a.nivel_id for a in solucion.asignaciones if a.profesor_id == p.id}) > 1
    )
    hoja = libro[multinivel.nombre]
    clases = [a for a in solucion.asignaciones if a.profesor_id == multinivel.id]
    # Cada clase es una celda combinada vertical (filas de 5 min: 45 o 50 min → 9 o 10 filas).
    alturas = sorted(r[2] - r[0] + 1 for r in _rangos(hoja) if r[1] == r[3] and r[1] > 1)
    assert len(alturas) == len(clases)
    assert set(alturas) <= {9, 10}
    assert hoja.row_dimensions[5].height == pytest.approx(6, abs=0.5)
    cuadricula = cuadricula_profesor(
        datos, generar_franjas(datos), solucion.asignaciones, multinivel.id
    )
    assert cuadricula.proporcional


def test_configuracion_de_impresion(libro: openpyxl.Workbook) -> None:
    for nombre in ("Resumen", "7°A", "CargaDocente"):
        hoja = libro[nombre]
        assert hoja.page_setup.orientation == "landscape"
        assert int(hoja.page_setup.paperSize) == 1  # carta
        assert hoja.sheet_properties.pageSetUpPr.fitToPage
        # 1 es el valor por defecto de OOXML; XlsxWriter lo omite (openpyxl lee None).
        assert (hoja.page_setup.fitToWidth or 1) == 1
        assert (hoja.page_setup.fitToHeight or 1) == 1
        assert hoja.page_margins.left == pytest.approx(0.25)
        assert "Colegio de ejemplo" in hoja.oddHeader.left.text
        assert "Página &P de &N" in hoja.oddFooter.right.text
        assert "Generado:" in hoja.oddFooter.left.text


def test_resumen_y_carga(libro: openpyxl.Workbook, resuelto: tuple[DatosEscuela, Solucion]) -> None:
    datos, solucion = resuelto
    resumen = libro["Resumen"]
    valores = {resumen.cell(r, 1).value: resumen.cell(r, 2).value for r in range(1, 10)}
    assert valores["Escuela"] == "Colegio de ejemplo"
    assert valores["Estado del solver"] == solucion.estado.value
    carga = libro["CargaDocente"]
    assert [c.value for c in carga[1]] == [
        "Profesor",
        "Minutos asignados",
        "Máximo",
        "Porcentaje",
        "Días con clase",
        "Huecos",
    ]
    assert carga.max_row == len(datos.profesores) + 1
    leyenda = libro["Leyenda"]
    assert leyenda.cell(2, 1).value == datos.asignaturas[0].nombre
    assert leyenda.cell(2, 2).fill.fgColor.rgb.endswith(datos.asignaturas[0].color[1:])


def test_opciones_de_inclusion(resuelto: tuple[DatosEscuela, Solucion], tmp_path: Path) -> None:
    datos, solucion = resuelto
    ruta = exportar_excel(
        datos,
        solucion,
        generar_franjas(datos),
        tmp_path / "solo_grupos.xlsx",
        OpcionesExportacion(profesores=False, espacios=False, resumen=False),
    )
    nombres = openpyxl.load_workbook(ruta).sheetnames
    assert nombres == [datos.nombre_grupo(g) for g in datos.grupos] + ["Leyenda"]


def test_nombres_de_hoja_validos() -> None:
    usados: set[str] = set()
    assert nombre_hoja("Ciencias: [Lab]/A*?", usados) == "Ciencias_ _Lab__A__"
    largo = "Profesora María de los Ángeles Fernández"
    primero = nombre_hoja(largo, usados)
    segundo = nombre_hoja(largo, usados)
    assert len(primero) == 31 and len(segundo) == 31 and primero != segundo
    assert segundo.endswith(" (2)")
    assert nombre_hoja("resumen", usados) == "resumen"
    assert nombre_hoja("RESUMEN", usados) == "RESUMEN (2)"


def test_cuadricula_grupo_intercala_recesos() -> None:
    datos = construir_datos(Tamano.PEQUENO)
    grupo = datos.grupos[0]  # K3A, preescolar
    cuadricula = cuadricula_grupo(datos, generar_franjas(datos), [], grupo.id)
    assert [f.receso for f in cuadricula.filas] == [None, None, None, "Recreo", None, None]
    assert cuadricula.filas[3].etiqueta == "10:15–10:45"


def test_cli_exportar(tmp_path: Path) -> None:
    from typer.testing import CliRunner

    from timetabling.cli import app

    ruta = tmp_path / "e.horario"
    runner = CliRunner()
    assert runner.invoke(app, ["ejemplo", str(ruta)]).exit_code == 0
    salida = tmp_path / "h.xlsx"
    r = runner.invoke(
        app, ["generar", str(ruta), "--tiempo-limite", "5", "--exportar", str(salida)]
    )
    assert r.exit_code == 0, r.output
    assert "Resumen" in openpyxl.load_workbook(salida).sheetnames
    r = runner.invoke(app, ["exportar", str(ruta), "--solucion", "99", "--salida", str(salida)])
    assert r.exit_code == 1 and "No existe la solución 99" in r.output
